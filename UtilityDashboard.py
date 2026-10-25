import io
import base64
import docx
import streamlit as st
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import pandas as pd
import plotly.graph_objects as go
from pydantic import BaseModel, Field
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.prompts import ChatPromptTemplate

# -----------------------------------------------------------------------------
# 1. PAGE CONFIGURATION & LOGIN LOGIC
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Utility Monitoring Dashboard",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="collapsed"
)

if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False

def login():
    # Hide the sidebar completely when on the login screen
    st.markdown(
        """
        <style>
            [data-testid="stSidebar"] {
                display: none;
            }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.title("🔒 Utility Dashboard Login")

    with st.form("login_form"):
        username = st.text_input("Username")
        password = st.text_input("Password", type="password")
        submit = st.form_submit_button("Sign In")

        if submit:
            clean_user = username.strip()
            clean_pass = password.strip()

            if clean_user == st.secrets["username"] and clean_pass == st.secrets["password"]:
                st.session_state["authenticated"] = True
                st.session_state["username"] = clean_user
                st.rerun()
            else:
                st.error("Invalid Username or Password")

if not st.session_state["authenticated"]:
    login()
    st.stop()


# -----------------------------------------------------------------------------
# 2. INITIALIZE SESSION STATE FOR ACCUMULATED DATAFRAMES
# -----------------------------------------------------------------------------
if "elec_history" not in st.session_state:
    st.session_state["elec_history"] = pd.DataFrame(columns=[
        "billing_period", "due_date", "total_amt_due", "actual_consumption",
        "effective_rate", "average_daily_consumption", "average_daily_cost", "average_temperature"
    ])

if "water_history" not in st.session_state:
    st.session_state["water_history"] = pd.DataFrame(columns=[
        "billing_period", "due_date", "total_amount_due", "actual_consumption", "effective_blended_rate"
    ])

if "solar_history" not in st.session_state:
    st.session_state["solar_history"] = pd.DataFrame(columns=[
        "timestamp", "daily_solar_production", "estimated_income",
        "cum_energy_month", "cum_energy_year", "cum_energy_lifetime",
        "cum_financial_month", "cum_financial_year", "cum_financial_lifetime"
    ])

# -----------------------------------------------------------------------------
# 3. SIDEBAR NAVIGATION & USER SESSION
# -----------------------------------------------------------------------------
st.sidebar.write(f"Logged in as: **{st.session_state['username']}**")
st.sidebar.markdown("---")
st.sidebar.markdown("### Navigation")

page = st.sidebar.radio(
    "Select Utility Page",
    ["💡 Electricity Agent", "☀️ Solar Agent", "💧 Water Agent", "📊 Combined Dashboard"],
    label_visibility="collapsed"
)

st.sidebar.markdown("---")
if st.sidebar.button("Logout"):
    st.session_state["authenticated"] = False
    st.rerun()


# -----------------------------------------------------------------------------
# 4. PYDANTIC SCHEMAS
# -----------------------------------------------------------------------------
class ElectricityMetrics(BaseModel):
    billing_period: str = Field(description="Billing period start and end dates")
    due_date: str = Field(description="Payment due date")
    total_amt_due: str = Field(description="Total amount due in PHP")
    actual_consumption: str = Field(description="Actual consumption in kWh")
    effective_rate: str = Field(description="Effective rate per kWh in PHP")
    average_daily_consumption: str = Field(description="Average daily consumption in kWh")
    average_daily_cost: str = Field(description="Average daily cost in PHP")
    average_temperature: str = Field(description="Average temperature recorded")


class SolarMetrics(BaseModel):
    timestamp: str = Field(description="Timestamp of the screenshot")
    daily_solar_production: str = Field(description="Daily solar production in kWh")
    estimated_income: str = Field(description="Estimated income today in PHP")
    cum_energy_month: str = Field(description="Cumulative energy generation for this month")
    cum_energy_year: str = Field(description="Cumulative energy generation for this year")
    cum_energy_lifetime: str = Field(description="Cumulative energy generation lifetime")
    cum_financial_month: str = Field(description="Cumulative financial return for this month")
    cum_financial_year: str = Field(description="Cumulative financial return for this year")
    cum_financial_lifetime: str = Field(description="Cumulative financial return lifetime")


class WaterMetrics(BaseModel):
    billing_period: str = Field(description="Billing period start and end dates")
    due_date: str = Field(description="Payment due date")
    total_amount_due: str = Field(description="Total amount due in PHP")
    actual_consumption: str = Field(description="Actual consumption in cubic meters (cu.m.)")
    effective_blended_rate: str = Field(description="Effective blended rate per cu.m. in PHP")


# -----------------------------------------------------------------------------
# 5. HELPER FUNCTIONS & GEMINI AGENTS
# -----------------------------------------------------------------------------
@st.cache_resource
def get_llm():
    return ChatGoogleGenerativeAI(model="gemini-2.5-flash", temperature=0)


llm = get_llm()


def process_file_bytes(file_bytes):
    return base64.b64encode(file_bytes).decode("utf-8")


def parse_electricity(file_bytes, mime_type):
    b64_str = process_file_bytes(file_bytes)
    prompt = ChatPromptTemplate.from_messages([
        ("system", "You are an expert utility auditing AI agent specialized in parsing electricity bills."),
        ("human", [
            {"type": "text", "text": "Extract all electricity metrics from this statement."},
            {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64_str}"}}
        ])
    ])
    structured_llm = llm.with_structured_output(ElectricityMetrics)
    return (prompt | structured_llm).invoke({})


def parse_solar(file_bytes, mime_type):
    b64_str = process_file_bytes(file_bytes)
    prompt = ChatPromptTemplate.from_messages([
        ("system", "You are an expert AI agent specialized in parsing solar dashboard screenshots."),
        ("human", [
            {"type": "text", "text": "Extract all solar production and return metrics from this image."},
            {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64_str}"}}
        ])
    ])
    structured_llm = llm.with_structured_output(SolarMetrics)
    return (prompt | structured_llm).invoke({})


def parse_water(file_bytes, mime_type):
    b64_str = process_file_bytes(file_bytes)
    prompt = ChatPromptTemplate.from_messages([
        ("system", "You are an expert AI agent specialized in parsing water utility bills."),
        ("human", [
            {"type": "text", "text": "Extract all water billing metrics from this statement."},
            {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64_str}"}}
        ])
    ])
    structured_llm = llm.with_structured_output(WaterMetrics)
    return (prompt | structured_llm).invoke({})


def generate_docx_report(elec_df, water_df, solar_df):
    """Generates a styled Word document (.docx) containing accumulated utility metrics."""
    doc = docx.Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_heading('Utility Monitoring & Performance Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(f"Generated on: {pd.Timestamp.now().strftime('%B %d, %Y - %I:%M %p')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.italic = True
    sub.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_df_table(heading_text, df):
        doc.add_heading(heading_text, level=1)
        if df.empty:
            p = doc.add_paragraph("No data recorded for this section.")
            p.runs[0].font.italic = True
            doc.add_paragraph()
            return

        table = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1])
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(df.columns):
            hdr_cells[col_idx].text = str(col_name).replace("_", " ").title()
            p = hdr_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(10)

        for row_idx, row in df.iterrows():
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, val in enumerate(row):
                row_cells[col_idx].text = str(val) if pd.notna(val) else "N/A"
                p = row_cells[col_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    add_df_table("💡 Electricity Statements (Meralco)", elec_df)
    add_df_table("💧 Water Statements (Maynilad)", water_df)
    add_df_table("☀️ Solar Production Logs", solar_df)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# -----------------------------------------------------------------------------
# 6. PAGE ROUTING & LAYOUTS
# -----------------------------------------------------------------------------

if page == "💡 Electricity Agent":
    st.title("💡 Electricity Bill Processing (Meralco)")
    st.write(
        "Upload your electricity bill statement (PDF or JPEG) to extract and append metrics to your master tracking table.")

    uploaded_file = st.file_uploader("Upload Electricity Bill", type=["pdf", "jpg", "jpeg"], key="elec")

    if uploaded_file:
        with st.spinner("AI Agent processing Electricity Bill..."):
            file_bytes = uploaded_file.read()
            mime_type = uploaded_file.type
            result = parse_electricity(file_bytes, mime_type)

            data_dict = result.model_dump()

            new_row_df = pd.DataFrame([data_dict])
            st.session_state["elec_history"] = pd.concat([st.session_state["elec_history"], new_row_df],
                                                         ignore_index=True)

            st.success("Extraction Complete & Added to Master History!")
            col1, col2 = st.columns([1, 1])

            with col1:
                st.markdown("### Latest Extracted Metrics")
                st.dataframe(new_row_df.T.reset_index().rename(columns={"index": "Metric", 0: "Value"}),
                             use_container_width=True, hide_index=True)

            with col2:
                st.markdown("### Key Metrics Overview")
                st.metric("Total Amount Due", data_dict.get("total_amt_due", "N/A"))
                st.metric("Actual Consumption", data_dict.get("actual_consumption", "N/A"))
                st.metric("Effective Rate", data_dict.get("effective_rate", "N/A"))

elif page == "☀️ Solar Agent":
    st.title("☀️ Solar Inverter Dashboard")
    st.write(
        "Upload your solar monitoring app screenshot (JPEG/PNG) to analyze and append metrics to your master tracking table.")

    uploaded_file = st.file_uploader("Upload Solar Screenshot", type=["jpg", "jpeg", "png"], key="solar")

    if uploaded_file:
        with st.spinner("AI Agent processing Solar Dashboard..."):
            file_bytes = uploaded_file.read()
            mime_type = uploaded_file.type
            result = parse_solar(file_bytes, mime_type)

            data_dict = result.model_dump()

            new_row_df = pd.DataFrame([data_dict])
            st.session_state["solar_history"] = pd.concat([st.session_state["solar_history"], new_row_df],
                                                          ignore_index=True)

            st.success("Extraction Complete & Added to Master History!")
            col1, col2 = st.columns([1, 1])

            with col1:
                st.markdown("### Latest Extracted Solar Metrics")
                st.dataframe(new_row_df.T.reset_index().rename(columns={"index": "Metric", 0: "Value"}),
                             use_container_width=True, hide_index=True)

            with col2:
                st.markdown("### Production & Returns")
                st.metric("Daily Production", data_dict.get("daily_solar_production", "N/A"))
                st.metric("Estimated Income Today", data_dict.get("estimated_income", "N/A"))
                st.metric("Lifetime Generation", data_dict.get("cum_energy_lifetime", "N/A"))

elif page == "💧 Water Agent":
    st.title("💧 Water Bill Processing (Maynilad)")
    st.write(
        "Upload your water bill statement (PDF or JPEG) to extract and append metrics to your master tracking table.")

    uploaded_file = st.file_uploader("Upload Water Bill", type=["pdf", "jpg", "jpeg"], key="water")

    if uploaded_file:
        with st.spinner("AI Agent processing Water Bill..."):
            file_bytes = uploaded_file.read()
            mime_type = uploaded_file.type
            result = parse_water(file_bytes, mime_type)

            data_dict = result.model_dump()

            new_row_df = pd.DataFrame([data_dict])
            st.session_state["water_history"] = pd.concat([st.session_state["water_history"], new_row_df],
                                                          ignore_index=True)

            st.success("Extraction Complete & Added to Master History!")
            col1, col2 = st.columns([1, 1])

            with col1:
                st.markdown("### Latest Extracted Water Metrics")
                st.dataframe(new_row_df.T.reset_index().rename(columns={"index": "Metric", 0: "Value"}),
                             use_container_width=True, hide_index=True)

            with col2:
                st.markdown("### Key Metrics Overview")
                st.metric("Total Amount Due", data_dict.get("total_amount_due", "N/A"))
                st.metric("Actual Consumption", data_dict.get("actual_consumption", "N/A"))
                st.metric("Blended Rate", data_dict.get("effective_blended_rate", "N/A"))

elif page == "📊 Combined Dashboard":
    st.title("📊 Master Utility Dashboard & History")
    st.write("Here is the accumulated historical database of all your uploaded utility statements and solar reports.")

    # Action buttons layout for exporting and clearing history
    col_btn1, col_btn2 = st.columns([1, 1])

    with col_btn1:
        docx_file = generate_docx_report(
            st.session_state["elec_history"],
            st.session_state["water_history"],
            st.session_state["solar_history"]
        )

        st.download_button(
            label="📄 Export Report to Word (.docx)",
            data=docx_file,
            file_name=f"Utility_Report_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )

    with col_btn2:
        if st.button("🗑️ Clear All Stored History", use_container_width=True):
            st.session_state["elec_history"] = st.session_state["elec_history"].iloc[0:0]
            st.session_state["water_history"] = st.session_state["water_history"].iloc[0:0]
            st.session_state["solar_history"] = st.session_state["solar_history"].iloc[0:0]
            st.rerun()

    st.markdown("---")

    st.subheader("💡 Electricity Bill History")
    if not st.session_state["elec_history"].empty:
        st.dataframe(st.session_state["elec_history"], use_container_width=True, hide_index=True)
    else:
        st.info("No electricity bills uploaded yet. Visit the Electricity Agent page to upload records.")

    st.markdown("---")

    st.subheader("💧 Water Bill History")
    if not st.session_state["water_history"].empty:
        st.dataframe(st.session_state["water_history"], use_container_width=True, hide_index=True)
    else:
        st.info("No water bills uploaded yet. Visit the Water Agent page to upload records.")

    st.markdown("---")

    st.subheader("☀️ Solar Production History")
    if not st.session_state["solar_history"].empty:
        st.dataframe(st.session_state["solar_history"], use_container_width=True, hide_index=True)
    else:
        st.info("No solar screenshots uploaded yet. Visit the Solar Agent page to upload records.")